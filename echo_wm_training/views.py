"""Views of the training clips: DICOM tags, a quick human view-labelling sheet, and the model's view head.

Echo files rarely say which view they are. Training needs to know which clips are apical 4-, 2- and 3-chamber:

  1. tags      : io.list_clips reads the view when the scanner wrote it (view_tag)
  2. labels    : a person labels a sample of clips (sheet below; about 2,000 clips, ~5 s each from a 4-frame strip)
  3. view head : train.py stage V learns views from 1 + 2; predict_views() then labels every clip, and only
                 confident predictions (config views.min_view_prob) are used for training

    python -m echo_wm_training.views sheet --clips /workspace/echo/clips.parquet --n 2000 --labellers ds1,ds2 --out /workspace/echo/views
    ... fill views_<code>.xlsx (view: A4C / A2C / A3C / other) ...
    python -m echo_wm_training.views read --dir /workspace/echo/views --out /workspace/echo/view_labels.csv

The sheet and its images stay inside MCP.
"""
from __future__ import annotations

import argparse
import struct
import zlib
from pathlib import Path
from typing import Callable, List, Optional

import numpy as np
import pandas as pd

from .runtime.segments import VIEW_CLASSES, VIEWS


def write_png(path: Path, img: np.ndarray) -> None:
    """Greyscale uint8 (H, W) -> PNG, no imaging library needed."""
    img = np.ascontiguousarray(img, dtype=np.uint8)
    h, w = img.shape
    raw = b"".join(b"\x00" + img[r].tobytes() for r in range(h))

    def chunk(t: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + t + data + struct.pack(">I", zlib.crc32(t + data) & 0xFFFFFFFF)
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 0, 0, 0, 0)) + \
        chunk(b"IDAT", zlib.compress(raw, 6)) + chunk(b"IEND", b"")
    Path(path).write_bytes(png)


def strip(frames: np.ndarray, n: int = 4) -> np.ndarray:
    """(T, H, W) float [0, 1] -> one image of n frames side by side (end-diastole to end-systole roughly covered)."""
    idx = np.linspace(0, len(frames) - 1, n).round().astype(int)
    return (np.concatenate([frames[i] for i in idx], axis=1) * 255).clip(0, 255).astype(np.uint8)


def export_view_sheet(clips: pd.DataFrame, out: str, labellers: List[str], n: int = 2000, seed: int = 20261017,
                      reader: Optional[Callable] = None, log: Callable[[str], None] = print) -> pd.DataFrame:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    from .io import read_clip
    from .runtime.preprocess import EchoInputError, prepare
    reader = reader or read_clip
    outp = Path(out)
    (outp / "img").mkdir(parents=True, exist_ok=True)
    pool = clips[clips.view_tag.isna()] if "view_tag" in clips and clips.view_tag.isna().sum() >= n else clips
    pick = pool.sample(n=min(n, len(pool)), random_state=seed).reset_index(drop=True)
    rows = []
    for k, r in pick.iterrows():
        cid = f"V{k + 1:05d}"
        try:
            p = prepare(reader(r.clip_uri))
        except (EchoInputError, Exception) as e:  # noqa: BLE001
            log(f"{cid}: skipped ({e})")
            continue
        write_png(outp / "img" / f"{cid}.png", strip(p.frames))
        rows.append({"clip_id": cid, "clip_uri": r.clip_uri, "study_id": r.study_id, "labeller": labellers[k % len(labellers)]})
    key = pd.DataFrame(rows)
    key.to_csv(outp / "view_key.csv", index=False)
    for who in labellers:
        part = key[key.labeller == who]
        wb = Workbook()
        ws = wb.active
        ws.title = "views"
        ws.append(["clip_id", "image", "view", "quality", "note"])
        for _, r in part.iterrows():
            ws.append([r.clip_id, f"img/{r.clip_id}.png", None, None, None])
        if len(part):
            dv = DataValidation(type="list", formula1='"' + ",".join(VIEW_CLASSES) + '"', allow_blank=True)
            dq = DataValidation(type="list", formula1='"good,limited,unusable"', allow_blank=True)
            ws.add_data_validation(dv)
            ws.add_data_validation(dq)
            dv.add(f"C2:C{ws.max_row}")
            dq.add(f"D2:D{ws.max_row}")
        rs = wb.create_sheet("how")
        for line in ("Open img/<clip_id>.png: 4 frames of the cine, apex at the top.",
                     "A4C: both ventricles and both atria, septum vertical in the middle.",
                     "A2C: LV and LA only (no RV, no aortic root); anterior wall on the right, inferior on the left.",
                     "A3C: LV, LA and the aortic root / LVOT (apical long axis).",
                     "other: parasternal, subcostal, suprasternal, off-axis or foreshortened apical views.",
                     "quality: unusable = the LV walls can't be seen; such clips are not used."):
            rs.append([line])
        wb.save(outp / f"views_{who}.xlsx")
    return key


def read_view_labels(d: str) -> pd.DataFrame:
    dp = Path(d)
    key = pd.read_csv(dp / "view_key.csv")
    parts = []
    for p in sorted(dp.glob("views_*.xlsx")):
        s = pd.read_excel(p, sheet_name="views", dtype=str)
        parts.append(s[["clip_id", "view", "quality"]])
    lab = pd.concat(parts) if parts else pd.DataFrame(columns=["clip_id", "view", "quality"])
    lab["view"] = lab.view.str.strip().str.upper().replace({"OTHER": "other"})
    bad = lab[lab.view.notna() & ~lab.view.isin(VIEW_CLASSES)]
    if len(bad):
        raise ValueError(f"{len(bad)} view labels not in {VIEW_CLASSES}, e.g. {bad.clip_id.head(3).tolist()}")
    out = key.merge(lab, on="clip_id", how="left")
    out.loc[out.quality.fillna("").str.lower().eq("unusable"), "view"] = "other"
    return out[out.view.notna()][["clip_uri", "study_id", "view", "quality"]]


def assign_views(clips: pd.DataFrame, labels: Optional[pd.DataFrame] = None, predicted: Optional[pd.DataFrame] = None,
                 min_prob: float = 0.9) -> pd.DataFrame:
    """view = human label, else DICOM tag, else a confident prediction; view_source says which."""
    c = clips.copy()
    c["view"], c["view_source"] = None, None
    if predicted is not None and len(predicted):
        p = predicted.set_index("clip_uri")
        ok = c.clip_uri.map(p.view_p).fillna(0) >= min_prob
        c.loc[ok, "view"] = c.loc[ok, "clip_uri"].map(p.view)
        c.loc[ok, "view_source"] = "model"
    if "view_tag" in c:
        t = c.view_tag.notna()
        c.loc[t, "view"], c.loc[t, "view_source"] = c.loc[t, "view_tag"], "tag"
    if labels is not None and len(labels):
        lab = labels.set_index("clip_uri").view
        h = c.clip_uri.isin(lab.index)
        c.loc[h, "view"], c.loc[h, "view_source"] = c.loc[h, "clip_uri"].map(lab), "label"
    c.loc[c.view.notna() & ~c.view.isin(VIEWS), "view"] = "other"
    return c


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sheet")
    s.add_argument("--clips", required=True)
    s.add_argument("--labellers", required=True)
    s.add_argument("--n", type=int, default=2000)
    s.add_argument("--out", required=True)
    r = sub.add_parser("read")
    r.add_argument("--dir", required=True)
    r.add_argument("--out", required=True)
    a = ap.parse_args()
    if a.cmd == "sheet":
        c = pd.read_parquet(a.clips) if a.clips.endswith(".parquet") else pd.read_csv(a.clips)
        k = export_view_sheet(c, a.out, [x.strip() for x in a.labellers.split(",") if x.strip()], a.n)
        print(f"{len(k)} clips exported to {a.out}")
    else:
        lab = read_view_labels(a.dir)
        lab.to_csv(a.out, index=False)
        print(lab.view.value_counts().to_string())


if __name__ == "__main__":
    main()
