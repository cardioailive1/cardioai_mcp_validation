"""Label audit: blinded sampling, auditor sheets, cardiologist adjudication, PPV / NPV / kappa, troponin
recomputation, RR irregularity for AF, and outline QC + grading."""
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
pytest.importorskip("openpyxl")

from annotation import label_audit as L  # noqa: E402
from annotation.select import load_tasks  # noqa: E402
from mri_training.dataset import load_config as mri_cfg, report_labels  # noqa: E402

CFG = load_tasks()["label_audit"]

# report template -> what a careful reader of the TEXT would enter (the truth for this test)
T = {
    "A": ("Subendocardial late gadolinium enhancement in the LAD territory consistent with prior myocardial infarction. "
          "Microvascular obstruction is present. LVEF 42 %. Infarct size 18 %.",
          dict(infarct=1, nonischemic_lge=0, lv_thrombus=0, lv_aneurysm=0, report_mvo=1, report_lvef=42, report_infarct_size=18)),
    "B": ("No late gadolinium enhancement. Normal biventricular size and function. LVEF 61 %.",
          dict(infarct=0, nonischemic_lge=0, lv_thrombus=0, lv_aneurysm=0, report_mvo=0, report_lvef=61)),
    "C": ("Mid-wall late gadolinium enhancement in the basal septum, consistent with myocarditis. LVEF 55 %.",
          dict(infarct=0, nonischemic_lge=1, lv_thrombus=0, lv_aneurysm=0, report_mvo=0, report_lvef=55)),
    "D": ("Transmural enhancement in the RCA territory consistent with infarct. Apical LV thrombus measuring 1 cm. "
          "Left ventricular aneurysm at the apex. Ejection fraction 30 %.",
          dict(infarct=1, nonischemic_lge=0, lv_thrombus=1, lv_aneurysm=1, report_mvo=0, report_lvef=30)),
    # a deliberately BAD rule (hedges ignored, see _manifest) calls this thrombus = 1; the text is uncertain ->
    # the audit must catch it. (The real rule masks it; this test proves the audit would catch the bug.)
    "E": ("Subendocardial late gadolinium enhancement in the LAD territory consistent with prior myocardial infarction. "
          "Possible small LV thrombus, cannot be excluded. LVEF 45 %.",
          dict(infarct=1, nonischemic_lge=0, lv_thrombus="U", lv_aneurysm=0, report_mvo=0, report_lvef=45)),
    "F": ("Cine imaging only; contrast was not given for this study. LVEF 58 %.",
          dict(infarct="NM", nonischemic_lge="NM", lv_thrombus=0, lv_aneurysm=0, report_mvo=0, report_lvef=58)),
}
MIX = {"A": 0.25, "B": 0.35, "C": 0.10, "D": 0.05, "E": 0.05, "F": 0.20}


def _manifest(n=1500, seed=0):
    rnd = np.random.default_rng(seed)
    kinds = rnd.choice(list(MIX), n, p=list(MIX.values()))
    cfg = mri_cfg()
    never = r"(?!x)x"                                   # planted bug: the label code ignores every hedge
    cfg["labels"]["context"] = {"pre_uncertain": never, "post_uncertain": never, "anywhere_uncertain": never}
    rows = []
    for i, k in enumerate(kinds):
        rows.append({"study_id": f"S{i:05d}", "person_id": i, "kind": k, "report_text": T[k][0], **report_labels(T[k][0], cfg)})
    return pd.DataFrame(rows)


def _fill(path, kinds, flip=()):
    """An auditor (or the cardiologist) who applies the rule sheet perfectly, except on the audit_ids in flip."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["audit"]
    head = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        aid = ws.cell(r, 1).value
        truth = T[kinds[aid]][1]
        for j, col in enumerate(head):
            if col in ("audit_id", "text", "evidence", "note"):
                continue
            v = truth.get(col)
            if aid in flip and col == "infarct" and v in (0, 1):
                v = 1 - v
            ws.cell(r, j + 1).value = v
    wb.save(path)


def _kinds(d, m):
    key = pd.read_csv(d / "key.csv", dtype=str)
    k = m.set_index("study_id").kind
    return {a: k[s] for a, s in zip(key.audit_id, key.study_id)}


def test_wilson_and_kappa():
    lo, hi = L.wilson(96, 100)
    assert 0.90 < lo < 0.91 and 0.98 < hi < 0.99
    assert L.cohen_kappa([1, 0, 1, 0], [1, 0, 1, 0]) == 1.0
    assert L.cohen_kappa([1, 1, 0, 0], [1, 0, 1, 0]) == 0.0
    assert L.norm_label("u") == "U" and L.norm_label(1.0) == 1.0 and L.norm_label("nm") == "NM" and L.norm_label("") is None
    with pytest.raises(ValueError):
        L.norm_label("maybe")
    assert L.norm_value("42%") == 42.0 and L.norm_value("-") is None


def test_sampling_is_blinded_deterministic_and_covers_every_label(tmp_path):
    m = _manifest()
    k1 = L.sample(m, "mri", ["ds1", "ds2"], CFG)
    k2 = L.sample(m, "mri", ["ds1", "ds2"], CFG)
    assert k1.equals(k2)
    assert k1.study_id.is_unique
    whys = k1.picked.str.split(";").explode()
    assert (whys == "random").sum() == CFG["mri"]["random"]
    for lab in CFG["mri"]["labels"]:
        n_pos = (k1.picked.str.contains(f"random|topup:{lab}") & (k1[f"auto_{lab}"] == 1)).sum()
        assert n_pos >= min(CFG["mri"]["min_positive"], int((m[lab] == 1).sum()))
    assert 0.15 < (k1.auditor_2 != "").mean() < 0.25 and (k1.auditor_1 != k1.auditor_2).all()
    L.write_auditor_sheets(k1, tmp_path, "mri", CFG)
    sheet = pd.read_excel(tmp_path / "sheet_ds1.xlsx", sheet_name="audit")
    assert not any(c.startswith("auto_") for c in sheet.columns) and "study_id" not in sheet.columns   # blinded
    assert "rules" in pd.ExcelFile(tmp_path / "sheet_ds1.xlsx").sheet_names
    with pytest.raises(FileExistsError):
        L.write_auditor_sheets(k1, tmp_path, "mri", CFG)


def test_full_cycle_catches_the_bad_thrombus_rule(tmp_path):
    m = _manifest()
    key = L.sample(m, "mri", ["ds1", "ds2"], CFG)
    L.write_auditor_sheets(key, tmp_path, "mri", CFG)
    kinds = _kinds(tmp_path, m)
    with pytest.raises(ValueError, match="not filled"):
        L.score(str(tmp_path), "mri", CFG, log=lambda s: None)
    double = key[key.auditor_2 != ""].audit_id.tolist()
    _fill(tmp_path / "sheet_ds1.xlsx", kinds)
    _fill(tmp_path / "sheet_ds2.xlsx", kinds, flip=set(double[:3]))          # three honest disagreements
    r1 = L.score(str(tmp_path), "mri", CFG, log=lambda s: None)
    assert r1["adjudication_sheet"].name == "adjudication_1.xlsx"
    res = r1["results"].set_index("column")
    assert res.loc["lv_thrombus", "pending_cardiologist"] > 0                 # every 'possible thrombus' went up
    assert res.loc["infarct", "pending_cardiologist"] >= 3                   # the auditor disagreements too
    adj = pd.read_excel(r1["adjudication_sheet"], sheet_name="audit")
    assert set(adj.audit_id) >= set(double[:3])
    assert not any(c.startswith("auto_") for c in adj.columns)
    # a second score before the cardiologist answers must not create another sheet
    assert L.score(str(tmp_path), "mri", CFG, log=lambda s: None)["adjudication_sheet"] is None
    _fill(r1["adjudication_sheet"], kinds)
    r2 = L.score(str(tmp_path), "mri", CFG, log=lambda s: None)
    res = r2["results"].set_index("column")
    assert (res.pending_cardiologist == 0).all()
    assert res.loc["lv_thrombus", "result"].startswith("FAIL") and res.loc["lv_thrombus", "ppv"] < 0.8
    for ok in ("infarct", "nonischemic_lge", "lv_aneurysm", "report_mvo", "report_lvef", "report_infarct_size"):
        assert res.loc[ok, "result"] == "PASS", (ok, res.loc[ok].to_dict())
    assert res.loc["infarct", "masked_but_stated"] == 0
    assert res.loc["infarct", "ds_vs_cardiologist"] == 1.0 and res.loc["infarct", "spot_n"] > 0
    assert 0 < res.loc["infarct", "kappa_ds1_ds2"] < 1
    errs = r2["errors"]
    assert set(errs.column) == {"lv_thrombus"} and set(errs.reference) == {"U"}
    assert {kinds[a] for a in errs.audit_id} == {"E"}


def test_troponin_recomputed_independently():
    ecg_cfg = {"labels": {"troponin_url_fallback": {"ng_per_l_male": 15, "ng_per_l_female": 10, "ng_per_l": 14, "ng_per_ml": 0.04},
                          "window_days": [-1, 2]}}
    e = pd.DataFrame({"ecg_id": ["1", "2", "3", "4"], "mi_dx": [1, 1, 0, 1], "stemi_dx": [1, 0, 0, 0], "nstemi_dx": [0, 1, 0, 0],
                      "mi_any_near": [1, 1, 0, 1], "trop_pos": [1, 1, 0, 1], "acute_mi": [1.0, 1.0, 0.0, 1.0],
                      "stemi": [1.0, 0.0, 0.0, 1.0]})
    rows = pd.DataFrame({
        "ecg_id": ["1", "2", "2", "3", "4"], "ecg_date": ["2022-01-01"] * 5, "gender_concept_id": [8507, 8532, 8532, 8507, 8507],
        "measurement_date": ["2022-01-02", "2022-01-01", "2022-01-09", "2022-01-01", "2022-01-01"],
        "value_as_number": [80.0, 12.0, 500.0, 14.0, 0.02], "unit_source_value": ["ng/L", "ng/L", "ng/L", "ng/L", "ng/mL"],
        "range_high": [np.nan] * 5})
    r = L.recompute_troponin(e, rows, ecg_cfg).set_index("ecg_id")
    assert r.trop_pos_recomputed.tolist() == [1, 1, 0, 0]       # female 12 > 10; day +8 is outside; ng/mL 0.02 < 0.04
    assert r.trop_pos_match.tolist() == [True, True, True, False]
    assert not r.loc["4", "acute_mi_match"] and not r.loc["4", "stemi_match"]
    assert set(r.attrs["units"].unit) == {"ng/l", "ng/ml"}
    sql = L.troponin_rows_sql({"ecg": {"table": "{work}.ecg", "columns": {"ecg_id": "ECG_ID", "acquisition_datetime": "ACQ",
                                                                           "person_id": "PID"}}}, ["a'1", "b"])
    assert "{{cdm}}.person" in sql and "{{cs:troponin}}" in sql and "'a''1'" in sql


def test_rr_irregularity_separates_af_from_sinus():
    import synthetic_ecg as E
    af = [L.rr_irregularity(E.ecg("af", seed=s), 500.0)["rr_cv"] for s in range(5)]
    sinus = [L.rr_irregularity(E.ecg("normal", seed=s), 500.0)["rr_cv"] for s in range(5)]
    assert min(af) > CFG["ecg"]["rr_irregular_cv"] and max(sinus) < CFG["ecg"]["rr_regular_cv"]
    m = pd.DataFrame({"ecg_id": ["a", "b"], "af": [1.0, 0.0], "waveform_uri": ["sinus", "af"]})
    fl = L.af_rhythm_flags(m, CFG, reader=lambda u: (E.ecg("normal" if u == "sinus" else "af", seed=1), 500.0, None, "mV"))
    assert fl.flag.tolist() == ["AF label but regular RR", "no AF label but irregular RR"]


def test_outline_qc_and_grading(tmp_path):
    torch = pytest.importorskip("torch")  # noqa: F841
    pytest.importorskip("nibabel")
    pytest.importorskip("pydicom")
    import synthetic_cmr as S
    from annotation.export import export_mri_study
    from test_annotation import _draw_truth
    root = tmp_path / "root"
    for i, sid in enumerate(["s1", "s2", "s3"]):
        lge, cine, truth = S.study(seed=10 + i, n_slices=7, phases=6)
        meta = export_mri_study(sid, lge, cine, root / "model" / sid, readers=["model"])
        if sid != "s3":                                    # s3 keeps the (empty) guess: fails the automatic checks
            _draw_truth(root / "model" / sid, truth, meta)
    man = pd.DataFrame({"study_id": ["s1", "s2", "s3"], "report_lvef": [np.nan, 5.0, np.nan],
                        "infarct": [1.0, 1.0, np.nan], "nonischemic_lge": [0.0, 0.0, np.nan], "report_mvo": [np.nan] * 3})
    t = L.outlines(str(root), str(tmp_path / "qc"), ["ds1", "ds2"], man, CFG).set_index("study_id")
    assert t.loc["s3", "auto_status"] == "rejected"
    assert "vs report 5%" in t.loc["s2", "flags"]
    assert t.loc["s1", "infarct_pct"] > 0
    from openpyxl import load_workbook
    for g in ("ds1", "ds2"):
        p = tmp_path / "qc" / f"grades_{g}.xlsx"
        wb = load_workbook(p)
        ws = wb["grades"]
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 5).value = "C" if ws.cell(r, 1).value == "s3" else "A"
        wb.save(p)
    out = L.score_outlines(str(tmp_path / "qc"), CFG, log=lambda s: None)
    g = out["grades"].set_index("study_id")
    assert g.loc["s3", "action"].startswith("redraw") and "graded C" in g.loc["s3", "cardiologist_review"]
    assert "graded A despite automatic flags" in g.loc["s2", "cardiologist_review"]
    assert out["summary"]["C"] == 1 and (tmp_path / "qc" / "cardiologist_review.csv").exists()


# --------------------------------------------------------------------------- automatic auditor
def test_auto_auditor_reads_like_the_rule_sheet():
    from annotation.auto_audit import AutoAuditor
    aa = AutoAuditor("mri", CFG)
    for k, (text, truth) in T.items():
        r = aa.read(text)
        assert {c: r[c] for c in truth} == truth, k
    r = aa.read("Small area of subendocardial enhancement, possibly representing a small infarct. LVEF 55-60%.")
    assert r["infarct"] == "U" and r["report_lvef"] == "U"                    # hedged sentence; range
    r = aa.read("Indication: rule out myocarditis. No late gadolinium enhancement. Left atrial appendage thrombus. "
                "LVEF 60 %. Echo LVEF 45%.")
    assert (r["infarct"], r["nonischemic_lge"], r["lv_thrombus"], r["report_lvef"]) == (0.0, 0.0, 0.0, 60.0)
    assert aa.read("No evidence of LV thrombus, aneurysm or microvascular obstruction.")["report_mvo"] == 0.0
    assert aa.read("LGE in the inferolateral wall.")["infarct"] == "U"          # enhancement, pattern not named
    e = AutoAuditor("ecg", CFG)
    assert [e.read(t)["af"] for t in ["Atrial fibrillation with RVR", "Sinus rhythm. History of atrial fibrillation",
                                      "Possible atrial flutter", "Sinus rhythm, no atrial fibrillation"]] == [1.0, 0.0, "U", 0.0]
    assert "thrombus" in aa.evidence(T["E"][0]) and "LVEF" in aa.evidence(T["E"][0])
    import inspect
    import annotation.auto_audit as A
    assert not re.search(r"^\s*(from|import) (mri|ecg)_training", inspect.getsource(A), re.M)   # independent of the training rules


def test_triage_settles_agreeing_rows_and_verifies_a_blinded_share(tmp_path):
    m = _manifest()
    key = L.sample(m, "mri", ["ds1", "ds2"], CFG, auto="triage")
    off = L.sample(m, "mri", ["ds1", "ds2"], CFG)
    assert key.audit_id.tolist() == off.audit_id.tolist()                  # same sample, fewer rows to read
    n_ds = (key.auditor_1 != "").sum()
    assert n_ds < 0.6 * len(key)
    settled = key[key.auto_settled]
    assert abs(settled.ds_verify.mean() - CFG["auto_audit"]["verify_fraction"]) < 0.02
    assert not key[key.picked.str.contains("hard:|missing:")].auto_settled.any()
    k = m.set_index("study_id").kind
    assert not any(k[s] == "E" for s in settled.study_id)                  # hedged thrombus never auto-settled
    L.write_auditor_sheets(key, tmp_path, "mri", CFG)
    sh = pd.read_excel(tmp_path / "sheet_ds1.xlsx", sheet_name="audit")
    assert "evidence" in sh.columns and not any(c.startswith(("aa_", "auto_")) for c in sh.columns)
    kinds = _kinds(tmp_path, m)
    _fill(tmp_path / "sheet_ds1.xlsx", kinds)
    _fill(tmp_path / "sheet_ds2.xlsx", kinds)
    r1 = L.score(str(tmp_path), "mri", CFG, log=lambda s: None)
    adj = pd.read_excel(r1["adjudication_sheet"], sheet_name="audit")
    assert set(adj.audit_id) & set(key[key.auto_settled & ~key.ds_verify].audit_id)   # spot-check reaches auto rows
    _fill(r1["adjudication_sheet"], kinds)
    res = L.score(str(tmp_path), "mri", CFG, log=lambda s: None)["results"].set_index("column")
    assert res.loc["lv_thrombus", "result"].startswith("FAIL")
    for ok in ("infarct", "nonischemic_lge", "lv_aneurysm", "report_mvo", "report_lvef"):
        assert res.loc[ok, "result"] == "PASS", (ok, res.loc[ok].to_dict())
        assert res.loc[ok, "auto_vs_reference"] == 1.0 and res.loc[ok, "auto_verify_n"] > 0
    assert res.loc["infarct", "auto_vs_cardiologist"] == 1.0


def test_triage_fails_when_the_automatic_auditor_is_unreliable(tmp_path):
    m = _manifest()
    key = L.sample(m, "mri", ["ds1", "ds2"], CFG, auto="triage")
    L.write_auditor_sheets(key, tmp_path, "mri", CFG)
    kinds = _kinds(tmp_path, m)
    for a in ("ds1", "ds2"):
        _fill(tmp_path / f"sheet_{a}.xlsx", kinds)
    k = pd.read_csv(tmp_path / "key.csv", dtype=str, keep_default_na=False)
    ver = k.index[k.ds_verify == "True"][:10]                               # pretend it misread 10 verified rows
    k.loc[ver, "aa_infarct"] = k.loc[ver, "aa_infarct"].map({"1.0": "0.0", "0.0": "1.0"}).fillna("U")
    k.to_csv(tmp_path / "key.csv", index=False)
    res = L.score(str(tmp_path), "mri", CFG, log=lambda s: None)["results"].set_index("column")
    assert "automatic auditor" in res.loc["infarct", "result"] and res.loc["infarct", "auto_vs_reference"] < 0.98


def test_assist_keeps_every_row_for_the_data_scientists(tmp_path):
    key = L.sample(_manifest(), "mri", ["ds1"], CFG, auto="assist")
    assert (key.auditor_1 == "ds1").all() and not key.auto_settled.any()
    L.write_auditor_sheets(key, tmp_path, "mri", CFG)
    sh = pd.read_excel(tmp_path / "sheet_ds1.xlsx", sheet_name="audit")
    assert list(sh.columns[:3]) == ["audit_id", "text", "evidence"] and sh.evidence.notna().mean() > 0.9
